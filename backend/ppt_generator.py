"""
ppt_generator.py – Generate a polished .pptx from the smart slide plan.

Slide types:
  title     – full-width title + subtitle + authors
  content   – header + full-width bullet list (font size adapts to bullet count)
  split     – header + bullets (left 55%) + figure (right 45%)
  figure    – header + centred full-width figure + caption
  table     – header + faithful native pptx table (proportional col widths,
               auto header-row detection, word-wrap, horizontal rules)
  thankyou  – centred closing slide
"""

import os
import re
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# ─────────────────────────────────────────────────────────────────────────────
# Theme Definitions
# ─────────────────────────────────────────────────────────────────────────────

class Theme:
    def __init__(self, name: str):
        self.name = name.lower()
        self.style = "classic" # Default structure

        if self.name == "light":
            self.style = "minimal"
            self.C_BG_DARK    = RGBColor(0xF8, 0xFA, 0xFC)
            self.C_BG_PANEL   = RGBColor(0xFF, 0xFF, 0xFF)
            self.C_BG_PANEL2  = RGBColor(0xF0, 0xF4, 0xF8)
            self.C_ACCENT     = RGBColor(0x00, 0x66, 0xCC)
            self.C_ACCENT2    = RGBColor(0x7C, 0x3A, 0xED)
            self.C_ACCENT3    = RGBColor(0x00, 0x66, 0xCC)
            self.C_TEXT_LIGHT = RGBColor(0x1A, 0x20, 0x2C)
            self.C_TEXT_DIM   = RGBColor(0x4A, 0x55, 0x68)
            self.C_WHITE      = RGBColor(0x00, 0x00, 0x00)
            
            # Table specific
            self.C_HDR_BG     = RGBColor(0xE2, 0xE8, 0xF0)
            self.C_ROW_ODD    = RGBColor(0xFA, 0xFA, 0xFA)
            self.C_ROW_EVEN   = RGBColor(0xF0, 0xF4, 0xF8)
            self.C_HDR_TXT    = RGBColor(0x1A, 0x20, 0x2C)
            self.C_ROW_TXT    = RGBColor(0x2D, 0x37, 0x48)
            self.C_RULE       = RGBColor(0x00, 0x66, 0xCC)
            self.C_TBL_BDR    = RGBColor(0xCB, 0xD5, 0xE0)
        elif self.name == "nature":
            self.style = "organic"
            self.C_BG_DARK    = RGBColor(0x0A, 0x1F, 0x15)
            self.C_BG_PANEL   = RGBColor(0x11, 0x2E, 0x1F)
            self.C_BG_PANEL2  = RGBColor(0x16, 0x3D, 0x29)
            self.C_ACCENT     = RGBColor(0xDF, 0xB7, 0x5D)
            self.C_ACCENT2    = RGBColor(0x4A, 0x90, 0xE2)
            self.C_ACCENT3    = RGBColor(0xA3, 0xD4, 0xB6)
            self.C_TEXT_LIGHT = RGBColor(0xF0, 0xF7, 0xF2)
            self.C_TEXT_DIM   = RGBColor(0x8B, 0xAA, 0x98)
            self.C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
            
            self.C_HDR_BG     = RGBColor(0x1B, 0x48, 0x31)
            self.C_ROW_ODD    = RGBColor(0x0A, 0x1F, 0x15)
            self.C_ROW_EVEN   = RGBColor(0x11, 0x2E, 0x1F)
            self.C_HDR_TXT    = RGBColor(0xFF, 0xFF, 0xFF)
            self.C_ROW_TXT    = RGBColor(0xF0, 0xF7, 0xF2)
            self.C_RULE       = RGBColor(0xDF, 0xB7, 0x5D)
            self.C_TBL_BDR    = RGBColor(0x27, 0x52, 0x3C)
        elif self.name == "sunset":
            self.style = "bold"
            self.C_BG_DARK    = RGBColor(0x2B, 0x0E, 0x18)
            self.C_BG_PANEL   = RGBColor(0x3B, 0x14, 0x23)
            self.C_BG_PANEL2  = RGBColor(0x4D, 0x1A, 0x2D)
            self.C_ACCENT     = RGBColor(0xFF, 0x7E, 0x67)
            self.C_ACCENT2    = RGBColor(0xFF, 0xC4, 0x78)
            self.C_ACCENT3    = RGBColor(0xFA, 0x9C, 0xA4)
            self.C_TEXT_LIGHT = RGBColor(0xFF, 0xF0, 0xF3)
            self.C_TEXT_DIM   = RGBColor(0xBA, 0x93, 0x9D)
            self.C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
            
            self.C_HDR_BG     = RGBColor(0x5E, 0x22, 0x38)
            self.C_ROW_ODD    = RGBColor(0x2B, 0x0E, 0x18)
            self.C_ROW_EVEN   = RGBColor(0x3B, 0x14, 0x23)
            self.C_HDR_TXT    = RGBColor(0xFF, 0xFF, 0xFF)
            self.C_ROW_TXT    = RGBColor(0xFF, 0xF0, 0xF3)
            self.C_RULE       = RGBColor(0xFF, 0x7E, 0x67)
            self.C_TBL_BDR    = RGBColor(0x73, 0x30, 0x48)
        else: # Default dark
            self.C_BG_DARK    = RGBColor(0x0D, 0x11, 0x2B)
            self.C_BG_PANEL   = RGBColor(0x16, 0x1C, 0x40)
            self.C_BG_PANEL2  = RGBColor(0x12, 0x18, 0x36)
            self.C_ACCENT     = RGBColor(0x00, 0xD4, 0xFF)
            self.C_ACCENT2    = RGBColor(0x7C, 0x3A, 0xED)
            self.C_ACCENT3    = RGBColor(0x00, 0xFF, 0xC8)
            self.C_TEXT_LIGHT = RGBColor(0xF0, 0xF4, 0xFF)
            self.C_TEXT_DIM   = RGBColor(0xA0, 0xAB, 0xC4)
            self.C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
            
            # Table specific
            self.C_HDR_BG     = RGBColor(0x00, 0x8B, 0xA8)
            self.C_ROW_ODD    = RGBColor(0x12, 0x18, 0x36)
            self.C_ROW_EVEN   = RGBColor(0x1A, 0x22, 0x4A)
            self.C_HDR_TXT    = RGBColor(0xFF, 0xFF, 0xFF)
            self.C_ROW_TXT    = RGBColor(0xF0, 0xF4, 0xFF)
            self.C_RULE       = RGBColor(0x00, 0xD4, 0xFF)
            self.C_TBL_BDR    = RGBColor(0x2A, 0x35, 0x60)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
FONT_BODY = "Calibri"

BULLET_FONT_SIZE = {"large": 26, "normal": 22, "small": 18}
BULLET_ROW_H     = {"large": 0.82, "normal": 0.68, "small": 0.56}


# ─────────────────────────────────────────────────────────────────────────────
# Low-level drawing helpers
# ─────────────────────────────────────────────────────────────────────────────

def _add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    return shape


def _add_textbox(slide, left, top, width, height, text,
                 font_size, bold=False, color=None,
                 align=PP_ALIGN.LEFT, word_wrap=True,
                 italic=False, font_name=FONT_BODY):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size   = Pt(font_size)
    run.font.bold   = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    run.font.name   = font_name
    return txBox


def _set_slide_background(slide, color: RGBColor):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def _draw_header(slide, title_text: str, theme: Theme, slide_num: str = ""):
    if theme.style == "minimal":
        _add_rect(slide, Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.02), theme.C_ACCENT)
        _add_textbox(slide, left=Inches(0.4), top=Inches(0.25), width=Inches(12.0), height=Inches(0.8), text=title_text, font_size=28, bold=True, color=theme.C_TEXT_LIGHT)
    elif theme.style == "organic":
        _add_rect(slide, Inches(0), Inches(0), Inches(0.2), Inches(1.2), theme.C_ACCENT)
        _add_rect(slide, Inches(0.2), Inches(0), SLIDE_W - Inches(0.2), Inches(1.2), theme.C_BG_PANEL)
        _add_textbox(slide, left=Inches(0.5), top=Inches(0.2), width=Inches(12.0), height=Inches(0.8), text=title_text, font_size=28, bold=True, color=theme.C_WHITE)
    elif theme.style == "bold":
        _add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.3), theme.C_BG_PANEL2)
        _add_rect(slide, Inches(0), Inches(1.2), SLIDE_W, Inches(0.1), theme.C_ACCENT)
        _add_textbox(slide, left=Inches(0.5), top=Inches(0.25), width=Inches(12.0), height=Inches(0.9), text=title_text, font_size=32, bold=True, color=theme.C_WHITE)
    else: # classic
        _add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.1), theme.C_BG_PANEL)
        _add_rect(slide, Inches(0), Inches(1.05), SLIDE_W, Inches(0.05), theme.C_ACCENT)
        _add_textbox(slide, left=Inches(0.35), top=Inches(0.15), width=Inches(12.0), height=Inches(0.82), text=title_text, font_size=28, bold=True, color=theme.C_WHITE)
        
    if slide_num:
        _add_textbox(
            slide,
            left=Inches(12.3), top=Inches(0.3),
            width=Inches(0.85), height=Inches(0.5),
            text=slide_num,
            font_size=11, color=theme.C_TEXT_DIM,
            align=PP_ALIGN.RIGHT,
        )


def _draw_bullet_column(slide, bullets: list[str], font_scale: str,
                        left_in: float, top_in: float, width_in: float, theme: Theme):
    fs = BULLET_FONT_SIZE.get(font_scale, 22)
    
    txBox = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(5.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(fs)
        p.font.name = FONT_BODY
        p.font.color.rgb = theme.C_TEXT_LIGHT
        p.space_after = Pt(fs * 0.8) # Space between paragraphs
        
        # Enable native bullets via oxml for perfect hanging indentation
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(int(0.35 * 914400)))    # Left margin
        pPr.set('indent', str(int(-0.35 * 914400))) # Hanging indent
        
        buChar = etree.SubElement(pPr, qn('a:buChar'))
        buChar.set('char', '▸')
        
        buClr = etree.SubElement(pPr, qn('a:buClr'))
        srgbClr = etree.SubElement(buClr, qn('a:srgbClr'))
        srgbClr.set('val', '{:02X}{:02X}{:02X}'.format(*theme.C_ACCENT3))


# ─────────────────────────────────────────────────────────────────────────────
# Slide builders
# ─────────────────────────────────────────────────────────────────────────────

def _build_title_slide(prs: Presentation, sd: dict, theme: Theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide, theme.C_BG_DARK)
    W, H = SLIDE_W, SLIDE_H

    if theme.style == "minimal":
        _add_rect(slide, Inches(0.4), Inches(0.4), Inches(12.5), Inches(6.7), theme.C_BG_PANEL)
        _add_rect(slide, Inches(0.4), Inches(0.4), Inches(12.5), Inches(0.05), theme.C_ACCENT)
        _add_textbox(slide, left=Inches(1.0), top=Inches(1.5), width=Inches(11.3), height=Inches(2.0), text=sd.get("title", "Presentation"), font_size=36, bold=True, color=theme.C_TEXT_LIGHT, align=PP_ALIGN.CENTER)
        if sd.get("subtitle"):
            _add_textbox(slide, left=Inches(1.0), top=Inches(3.8), width=Inches(11.3), height=Inches(1.0), text=sd["subtitle"], font_size=20, color=theme.C_ACCENT, align=PP_ALIGN.CENTER)
        if sd.get("authors"):
            _add_textbox(slide, left=Inches(1.0), top=Inches(6.0), width=Inches(11.3), height=Inches(0.6), text=sd["authors"], font_size=14, color=theme.C_TEXT_DIM, align=PP_ALIGN.CENTER)
    elif theme.style == "organic":
        _add_rect(slide, Inches(0), Inches(0), W, Inches(0.4), theme.C_ACCENT)
        _add_rect(slide, Inches(0), H - Inches(0.4), W, Inches(0.4), theme.C_ACCENT)
        _add_textbox(slide, left=Inches(1.0), top=Inches(1.5), width=Inches(11.3), height=Inches(2.0), text=sd.get("title", "Presentation"), font_size=36, bold=True, color=theme.C_WHITE, align=PP_ALIGN.CENTER)
        if sd.get("subtitle"):
            _add_textbox(slide, left=Inches(1.0), top=Inches(3.8), width=Inches(11.3), height=Inches(1.0), text=sd["subtitle"], font_size=20, color=theme.C_ACCENT, align=PP_ALIGN.CENTER)
        if sd.get("authors"):
            _add_textbox(slide, left=Inches(1.0), top=Inches(6.0), width=Inches(11.3), height=Inches(0.6), text=sd["authors"], font_size=14, color=theme.C_TEXT_DIM, align=PP_ALIGN.CENTER)
    elif theme.style == "bold":
        _add_rect(slide, Inches(0), Inches(0), W, Inches(4.5), theme.C_BG_PANEL)
        _add_rect(slide, Inches(0), Inches(4.5), W, Inches(0.15), theme.C_ACCENT)
        _add_textbox(slide, left=Inches(0.5), top=Inches(0.8), width=Inches(12.3), height=Inches(3.0), text=sd.get("title", "Presentation"), font_size=42, bold=True, color=theme.C_WHITE, align=PP_ALIGN.LEFT)
        if sd.get("subtitle"):
            _add_textbox(slide, left=Inches(0.5), top=Inches(5.0), width=Inches(12.3), height=Inches(1.0), text=sd["subtitle"], font_size=22, color=theme.C_ACCENT, align=PP_ALIGN.LEFT)
        if sd.get("authors"):
            _add_textbox(slide, left=Inches(0.5), top=Inches(6.5), width=Inches(12.3), height=Inches(0.6), text=sd["authors"], font_size=14, color=theme.C_TEXT_DIM, align=PP_ALIGN.LEFT)
    else: # classic
        _add_rect(slide, Inches(0), Inches(0), Inches(0.08), H, theme.C_ACCENT)
        _add_textbox(
            slide, left=Inches(0.4), top=Inches(1.2),
            width=Inches(8.5), height=Inches(2.2),
            text=sd.get("title", "Presentation"), font_size=34, bold=True, color=theme.C_WHITE,
        )
        _add_rect(slide, Inches(0.4), Inches(3.55), Inches(7.5), Inches(0.05), theme.C_ACCENT)
        if sd.get("subtitle"):
            _add_textbox(
                slide, left=Inches(0.4), top=Inches(3.7),
                width=Inches(8.5), height=Inches(0.75),
                text=sd["subtitle"], font_size=18, color=theme.C_ACCENT,
            )
        if sd.get("motivation"):
            _add_textbox(
                slide, left=Inches(0.4), top=Inches(4.6),
                width=Inches(8.0), height=Inches(0.65),
                text=f"💡  {sd['motivation']}", font_size=13, color=theme.C_TEXT_DIM,
            )
        if sd.get("authors"):
            _add_textbox(
                slide, left=Inches(0.4), top=Inches(6.55),
                width=Inches(9.0), height=Inches(0.6),
                text=sd["authors"], font_size=12, color=theme.C_TEXT_DIM,
            )
        _add_rect(slide, Inches(9.3), Inches(0), Inches(4.03), H, theme.C_BG_PANEL)
        _add_rect(slide, Inches(9.3), Inches(0), Inches(0.05), H, theme.C_ACCENT2)
        _add_textbox(
            slide, left=Inches(9.5), top=Inches(3.1),
            width=Inches(3.5), height=Inches(1.0),
            text="Research Presentation",
            font_size=14, bold=True, color=theme.C_ACCENT2, align=PP_ALIGN.CENTER,
        )


def _build_content_slide(prs: Presentation, sd: dict, slide_num: str, theme: Theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide, theme.C_BG_DARK)
    _draw_header(slide, sd.get("title", "Overview"), theme, slide_num)
    _add_rect(slide, Inches(0.35), Inches(1.25), Inches(12.6), Inches(5.85), theme.C_BG_PANEL)
    _draw_bullet_column(slide, sd.get("bullets", []),
                        sd.get("font_scale", "normal"),
                        left_in=0.5, top_in=1.5, width_in=12.4, theme=theme)


def _build_split_slide(prs: Presentation, sd: dict, slide_num: str, theme: Theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide, theme.C_BG_DARK)
    _draw_header(slide, sd.get("title", "Details"), theme, slide_num)

    _add_rect(slide, Inches(0.35), Inches(1.25), Inches(7.0), Inches(5.85), theme.C_BG_PANEL)
    _add_rect(slide, Inches(7.4), Inches(1.3), Inches(0.04), Inches(5.7), theme.C_ACCENT2)
    _add_rect(slide, Inches(7.5), Inches(1.25), Inches(5.5), Inches(5.85), theme.C_BG_PANEL2)

    _draw_bullet_column(slide, sd.get("bullets", []),
                        sd.get("font_scale", "normal"),
                        left_in=0.5, top_in=1.5, width_in=6.8, theme=theme)

    img_path = sd.get("image_path")
    if img_path and os.path.exists(img_path):
        try:
            slide.shapes.add_picture(
                img_path,
                left=Inches(7.55), top=Inches(1.4),
                width=Inches(5.35), height=Inches(5.3),
            )
        except Exception:
            pass

    caption = sd.get("caption", "")
    if caption:
        _add_textbox(
            slide, left=Inches(7.5), top=Inches(6.85),
            width=Inches(5.5), height=Inches(0.45),
            text=caption[:100], font_size=10, color=theme.C_TEXT_DIM,
            align=PP_ALIGN.CENTER, italic=True,
        )


def _build_figure_slide(prs: Presentation, sd: dict, slide_num: str, theme: Theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide, theme.C_BG_DARK)
    _draw_header(slide, sd.get("title", "Figure"), theme, slide_num)

    img_path = sd.get("image_path")
    if img_path and os.path.exists(img_path):
        try:
            slide.shapes.add_picture(
                img_path,
                left=Inches(1.5), top=Inches(1.3),
                width=Inches(10.3), height=Inches(5.5),
            )
        except Exception:
            pass

    caption = sd.get("caption", "")
    if caption:
        _add_textbox(
            slide, left=Inches(0.5), top=Inches(6.9),
            width=Inches(12.3), height=Inches(0.45),
            text=caption, font_size=11, color=theme.C_TEXT_DIM,
            align=PP_ALIGN.CENTER, italic=True,
        )


def _build_table_slide(prs: Presentation, sd: dict, slide_num: str, theme: Theme):
    """
    Render a native python-pptx table that faithfully replicates the paper's table.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide, theme.C_BG_DARK)

    # Use caption as header if title isn't very descriptive
    header_title = sd.get("title", "Table")
    if header_title.lower() == "table" and sd.get("table_caption"):
        header_title = sd.get("table_caption")
    _draw_header(slide, header_title[:80], theme, slide_num)

    rows_data: list[list[str]] = sd.get("table_rows", [])
    col_widths_prop: list[float] = sd.get("col_widths", [])

    # We will also add the full caption at the bottom of the slide
    caption = sd.get("table_caption", "") or sd.get("caption", "")

    # ── Fallback: no structured rows ─────────────────────────────────────────
    if not rows_data:
        _add_rect(slide, Inches(0.35), Inches(1.25), Inches(12.6), Inches(5.85), theme.C_BG_PANEL)
        raw_text = sd.get("table_text", "(No table data extracted from PDF)")
        lines = raw_text.replace("\r", "").split("\n")
        wrapped = []
        for line in lines:
            wrapped.extend(textwrap.wrap(line, 90) if len(line) > 90 else [line])
        _add_textbox(
            slide, left=Inches(0.55), top=Inches(1.45),
            width=Inches(12.2), height=Inches(5.5),
            text="\n".join(wrapped[:30]),
            font_size=12, color=theme.C_TEXT_LIGHT, font_name="Courier New",
        )
        return

    # ── Cap to slide capacity ─────────────────────────────────────────────────
    MAX_ROWS = 18
    MAX_COLS = 9

    rows_data = rows_data[:MAX_ROWS]
    n_cols = min(max(len(r) for r in rows_data), MAX_COLS)

    rows_data = [
        [(r[c] if c < len(r) else "") for c in range(n_cols)]
        for r in rows_data
    ]
    n_rows = len(rows_data)

    # ── Table layout ──────────────────────────────────────────────────────────
    tbl_left   = Inches(0.4)
    tbl_top    = Inches(1.3)
    tbl_width  = Inches(12.5)
    tbl_height = Inches(5.4) # Slightly shorter to fit caption at bottom

    if col_widths_prop and len(col_widths_prop) >= n_cols:
        ratios = list(col_widths_prop[:n_cols])
    else:
        col_max_len = [1] * n_cols
        for row in rows_data:
            for ci, cell in enumerate(row):
                col_max_len[ci] = max(col_max_len[ci], len(str(cell)))
        total_len = sum(col_max_len) or 1
        ratios = [l / total_len for l in col_max_len]

    ratio_sum = sum(ratios) or 1.0
    ratios = [r / ratio_sum for r in ratios]
    col_widths_emu = [int(tbl_width * r) for r in ratios]
    col_widths_emu[-1] += int(tbl_width) - sum(col_widths_emu)

    if n_rows <= 6:    font_data, font_hdr = 12, 12
    elif n_rows <= 10: font_data, font_hdr = 11, 11
    elif n_rows <= 14: font_data, font_hdr = 10, 10
    else:              font_data, font_hdr = 9,  9

    def _is_header_row(row: list[str]) -> bool:
        non_empty = [c for c in row if c.strip()]
        if not non_empty:
            return True
        numeric = sum(1 for c in non_empty if re.match(r'^[-+]?[\d.,\s%±]+$', c))
        return numeric < len(non_empty) / 2

    n_header_rows = 0
    for ri in range(min(4, n_rows)):
        if _is_header_row(rows_data[ri]):
            n_header_rows = ri + 1
        else:
            break
    if n_header_rows == 0:
        n_header_rows = 1

    tbl_shape = slide.shapes.add_table(
        n_rows, n_cols, tbl_left, tbl_top, tbl_width, tbl_height,
    ).table

    for ci, cw in enumerate(col_widths_emu):
        tbl_shape.columns[ci].width = cw
    for ri in range(n_rows):
        tbl_shape.rows[ri].height = Inches(0.3)

    for ri, row in enumerate(rows_data):
        is_hdr = ri < n_header_rows
        bg     = theme.C_HDR_BG  if is_hdr else (theme.C_ROW_ODD if ri % 2 == 1 else theme.C_ROW_EVEN)
        fg     = theme.C_HDR_TXT if is_hdr else theme.C_ROW_TXT
        fsz    = font_hdr  if is_hdr else font_data

        for ci, cell_text in enumerate(row):
            cell = tbl_shape.cell(ri, ci)
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            run = para.add_run()
            run.text = str(cell_text)
            run.font.size  = Pt(fsz)
            run.font.bold  = is_hdr
            run.font.name  = FONT_BODY
            run.font.color.rgb = fg

            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("a:solidFill")):
                tcPr.remove(old)
            solidFill = etree.SubElement(tcPr, qn("a:solidFill"))
            srgbClr   = etree.SubElement(solidFill, qn("a:srgbClr"))
            srgbClr.set("val", "{:02X}{:02X}{:02X}".format(*bg))

            def _set_border(tcPr_el, tag, width_pt, rgb_color):
                for old in tcPr_el.findall(qn(tag)):
                    tcPr_el.remove(old)
                ln = etree.SubElement(tcPr_el, qn(tag))
                if rgb_color is None:
                    ln.set("w", "0")
                    etree.SubElement(ln, qn("a:noFill"))
                else:
                    ln.set("w", str(int(width_pt * 12700)))
                    sfill = etree.SubElement(ln, qn("a:solidFill"))
                    sc    = etree.SubElement(sfill, qn("a:srgbClr"))
                    sc.set("val", "{:02X}{:02X}{:02X}".format(*rgb_color))

            is_last_hdr = (ri == n_header_rows - 1)
            bot_color   = theme.C_RULE if is_last_hdr else theme.C_TBL_BDR
            bot_w       = 1.5    if is_last_hdr else 0.5

            _set_border(tcPr, "a:lnT", 0.5, theme.C_TBL_BDR)
            _set_border(tcPr, "a:lnB", bot_w, bot_color)
            _set_border(tcPr, "a:lnL", 0, None)
            _set_border(tcPr, "a:lnR", 0, None)

    # Add Caption at bottom
    if caption:
        _add_textbox(
            slide, left=Inches(0.4), top=Inches(6.8),
            width=Inches(12.5), height=Inches(0.6),
            text=caption, font_size=10, color=theme.C_TEXT_DIM,
            align=PP_ALIGN.LEFT, italic=True,
        )


def _build_thankyou_slide(prs: Presentation, sd: dict, theme: Theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide, theme.C_BG_DARK)
    W, H = SLIDE_W, SLIDE_H

    _add_rect(slide, Inches(3.2), Inches(1.8), Inches(6.9), Inches(3.8), theme.C_BG_PANEL)
    _add_rect(slide, Inches(3.2), Inches(1.8), Inches(0.07), Inches(3.8), theme.C_ACCENT)
    _add_rect(slide, Inches(3.2), Inches(5.55), Inches(6.9), Inches(0.05), theme.C_ACCENT)

    _add_textbox(
        slide, left=Inches(0), top=Inches(2.3),
        width=W, height=Inches(1.4),
        text="Thank You", font_size=48, bold=True, color=theme.C_WHITE,
        align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        slide, left=Inches(0), top=Inches(3.9),
        width=W, height=Inches(0.75),
        text=sd.get("subtitle", "Questions & Discussion"),
        font_size=22, color=theme.C_ACCENT, align=PP_ALIGN.CENTER,
    )
    if sd.get("authors"):
        _add_textbox(
            slide, left=Inches(0), top=Inches(6.55),
            width=W, height=Inches(0.6),
            text=sd["authors"], font_size=12, color=theme.C_TEXT_DIM,
            align=PP_ALIGN.CENTER,
        )


# ─────────────────────────────────────────────────────────────────────────────
# Main generator
# ─────────────────────────────────────────────────────────────────────────────

def generate_pptx(slide_plan: list[dict], output_path: str, theme_name: str = "dark") -> str:
    """Build the full .pptx and save it. Returns the output path."""
    theme = Theme(theme_name)
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    numbered_types = {"content", "split", "figure", "table"}
    total_numbered = sum(1 for s in slide_plan if s["type"] in numbered_types)
    counter = 0

    for sd in slide_plan:
        stype = sd["type"]

        if stype == "title":
            _build_title_slide(prs, sd, theme)

        elif stype == "content":
            counter += 1
            _build_content_slide(prs, sd, f"{counter}/{total_numbered}", theme)

        elif stype == "split":
            counter += 1
            _build_split_slide(prs, sd, f"{counter}/{total_numbered}", theme)

        elif stype == "figure":
            counter += 1
            _build_figure_slide(prs, sd, f"{counter}/{total_numbered}", theme)

        elif stype == "table":
            counter += 1
            _build_table_slide(prs, sd, f"{counter}/{total_numbered}", theme)

        elif stype == "thankyou":
            _build_thankyou_slide(prs, sd, theme)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    return output_path
